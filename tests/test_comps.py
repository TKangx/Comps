import math
import sys
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from comps.analysis import implied_price, percentile, run_comps  # noqa: E402
from comps.data import CompanyData, YahooProvider  # noqa: E402
from comps.excel import build_workbook  # noqa: E402
from comps.multiples import MULTIPLES, classify, select_multiples  # noqa: E402
from comps.peers import find_peers, score_peer  # noqa: E402
from comps.pipeline import analyze  # noqa: E402
from comps.sample import SAMPLE, SampleProvider  # noqa: E402


def co(**kw):
    base = dict(ticker="T", name="T", sector="Technology", industry="Software", price=10.0,
                market_cap=1000.0, total_debt=200.0, cash=100.0, revenue=500.0, ebitda=100.0,
                ebit=80.0, net_income=50.0)
    base.update(kw)
    return CompanyData(**base)


def test_enterprise_value_and_shares():
    c = co()
    assert c.enterprise_value == 1100.0
    assert c.net_debt == 100.0
    assert c.shares == 100.0


def test_multiple_math_and_nm():
    c = co()
    assert MULTIPLES["ev_ebitda"].value(c) == pytest.approx(11.0)
    assert MULTIPLES["pe_ltm"].value(c) == pytest.approx(20.0)
    assert MULTIPLES["pe_ltm"].value(co(net_income=-5)) is None        # negative -> NM
    assert MULTIPLES["pe_ltm"].value(co(net_income=1)) is None         # 1000x > cap -> NM
    assert MULTIPLES["ev_ntm_revenue"].value(c) is None                # missing -> NM


def test_percentile_matches_excel_percentile_inc():
    vals = [11.2, 7.6, 33.1, 13.0, 31.9]
    assert percentile(vals, 0.25) == pytest.approx(11.2)
    assert percentile(vals, 0.75) == pytest.approx(31.9)
    assert percentile([1, 2, 3, 4], 0.25) == pytest.approx(1.75)


def test_implied_price_ev_vs_equity():
    t = co()
    # EV multiple: (12 * 100 - net debt 100) / 100 shares = 11
    assert implied_price(t, "ev_ebitda", 12.0) == pytest.approx(11.0)
    # Equity multiple: 25 * 50 / 100 = 12.5
    assert implied_price(t, "pe_ltm", 25.0) == pytest.approx(12.5)
    assert implied_price(t, "pe_ltm", None) is None


@pytest.mark.parametrize("kw,profile", [
    (dict(sector="Financial Services", industry="Banks - Regional"), "financial"),
    (dict(sector="Financial Services", industry="Credit Services"), "mature"),
    (dict(sector="Real Estate", industry="REIT - Retail"), "reit"),
    (dict(ebitda=-10, net_income=-20), "unprofitable"),
    (dict(revenue_growth=0.35), "high_growth"),
    (dict(sector="Energy", industry="Oil & Gas E&P"), "capital_intensive"),
    (dict(), "mature"),
])
def test_classify(kw, profile):
    assert classify(co(**kw)) == profile


def test_bank_never_gets_ev_multiples():
    bank = co(sector="Financial Services", industry="Banks - Diversified", book_value=800, forward_earnings=60)
    peers = [co(ticker=f"P{i}", sector="Financial Services", industry="Banks - Diversified",
                book_value=700 + i * 50, forward_earnings=55 + i) for i in range(4)]
    sel = select_multiples(bank, peers)
    assert sel.multiples and all(MULTIPLES[k].numerator == "equity" for k in sel.multiples)


def test_peer_scoring_prefers_same_industry_and_size():
    t = co(revenue_growth=0.1, ebitda_margin=0.2, gross_margin=0.5)
    close = co(ticker="A", revenue_growth=0.1, ebitda_margin=0.2, gross_margin=0.5)
    far = co(ticker="B", sector="Energy", industry="Oil", market_cap=1.0, revenue_growth=0.5,
             ebitda_margin=0.6, gross_margin=0.1)
    assert score_peer(t, close).score == pytest.approx(90.0)  # everything but the investor-grouping bonus
    assert score_peer(t, close, market_peer=True).score == pytest.approx(100.0)
    assert score_peer(t, far).score < 20


def test_find_peers_filters_and_dedupes():
    p = SampleProvider()
    ranked = find_peers(p, p.get_company("AMZN"))
    tickers = [s.ticker for s in ranked]
    assert "AMZN" not in tickers
    assert "SAP.DE" not in tickers                     # foreign line dropped for US target
    assert not ({"GOOG", "GOOGL"} <= set(tickers))     # one share class only
    assert ranked == sorted(ranked, key=lambda s: -s.score)
    assert tickers[0] == "BABA"


def test_end_to_end_sample():
    a = analyze("AMZN", provider=SampleProvider(), n_peers=5)
    r = a.result
    assert len(r.peers) == 5
    assert r.selection.profile == "mature"
    assert r.blended_price and r.blended_price > 0
    assert "AMZN" in r.verdict
    ev = next(m for m in r.multiples if m.key == "ev_ebitda")
    assert ev.median == pytest.approx(13.0303, rel=1e-3)


def test_manual_peers_respected():
    a = analyze("AMZN", provider=SampleProvider(), manual_peers=["MSFT", "GOOGL", "NOPE"])
    assert [p.ticker for p in a.result.peers] == ["MSFT", "GOOGL"]


def test_excel_structure_and_formulas(tmp_path):
    a = analyze("AMZN", provider=SampleProvider())
    wb = build_workbook(a.result, a.candidates)
    assert wb.sheetnames == ["Dashboard", "Comps", "Valuation", "Peer Screen", "Lists", "Notes"]
    path = tmp_path / "out.xlsx"
    wb.save(path)
    ws = load_workbook(path)["Comps"]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("MEDIAN(" in f for f in formulas)
    assert any("PERCENTILE(" in f for f in formulas)
    # target is row 6, peer stats exclude it (start at row 7)
    assert any("MEDIAN(" in f and "7:" in f for f in formulas)
    assert ws["A6"].value == "AMZN" and ws["A7"].value == "BABA"


# ---------------------------------------------------------------- Yahoo parsing
class FakeTicker:
    def __init__(self, info, est=None):
        self.info = info
        self.revenue_estimate = est if est is not None else pd.DataFrame()
        self.fast_info = {"last_price": 0.14}  # CNY -> USD


def test_yahoo_provider_parses_and_converts_currency(monkeypatch):
    info = {
        "longName": "Foreign Co", "sector": "Consumer Cyclical", "industry": "Internet Retail",
        "currency": "USD", "financialCurrency": "CNY", "quoteType": "EQUITY",
        "currentPrice": 100.0, "marketCap": 200e9, "totalDebt": 100e9, "totalCash": 300e9,
        "totalRevenue": 1000e9, "ebitda": 150e9, "operatingMargins": 0.12, "netIncomeToCommon": 100e9,
        "bookValue": 500.0, "forwardPE": 10.0, "dividendYield": 0.5, "revenueGrowth": 0.08,
    }
    est = pd.DataFrame({"avg": [1100e9, 1200e9]}, index=["0y", "+1y"])
    fake_yf = SimpleNamespace(Ticker=lambda sym: FakeTicker(info, est) if sym == "FOO"
                              else FakeTicker({}, None))
    prov = YahooProvider.__new__(YahooProvider)
    prov.yf, prov._fx_cache = fake_yf, {}
    c = prov.get_company("foo")
    assert c.ticker == "FOO"
    assert c.revenue == pytest.approx(140e9)               # converted CNY -> USD
    assert c.ebit == pytest.approx(140e9 * 0.12)
    assert c.enterprise_value == pytest.approx(200e9 + 14e9 - 42e9)
    assert c.book_value == pytest.approx(500 * 0.14 * 2e9)  # per-share BV x implied shares
    assert c.forward_earnings == pytest.approx(20e9)
    assert c.dividend_yield == pytest.approx(0.005)
    assert 1100e9 * 0.14 <= c.ntm_revenue <= 1200e9 * 0.14


def test_yahoo_provider_pence_listing():
    info = {"currency": "GBp", "financialCurrency": "GBP", "currentPrice": 250.0, "marketCap": 500e9,
            "totalRevenue": 2e9}
    prov = YahooProvider.__new__(YahooProvider)
    prov.yf, prov._fx_cache = SimpleNamespace(Ticker=lambda s: FakeTicker(info)), {}
    c = prov.get_company("X.L")
    assert c.currency == "GBP" and c.price == pytest.approx(2.5) and c.market_cap == pytest.approx(5e9)
