"""The Excel template + the xlwings refresh path (with a fake xlwings Book)."""
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from comps import excel  # noqa: E402
from comps.data import DataError, resolve_ticker  # noqa: E402
from comps.sample import SampleProvider  # noqa: E402

RECALC = Path("/root/.claude/skills").glob("**/xlsx/scripts/recalc.py")


class FakeRange:
    """Mimics xlwings: assigning a scalar or 2D list writes from the range's top-left cell."""

    def __init__(self, wb, name):
        self.wb, self.name = wb, name

    @property
    def value(self):
        ws, first = excel._anchor(self.wb, self.name)
        return ws[first].value

    @value.setter
    def value(self, v):
        excel.fill(self.wb, {self.name: v if isinstance(v, list) else [[v]]})


class FakeBook:
    def __init__(self, wb):
        self.wb = wb
        self.names = {n: type("N", (), {"refers_to_range": FakeRange(wb, n)})() for n in excel.NAMES}


def test_resolve_ticker_by_name_and_symbol():
    p = SampleProvider()
    assert resolve_ticker(p, "AMZN") == "AMZN"
    assert resolve_ticker(p, "amazon") == "AMZN"
    assert resolve_ticker(p, "Alibaba") == "BABA"
    with pytest.raises(DataError):
        resolve_ticker(p, "Not A Real Company")


def test_template_names_point_at_real_cells():
    wb = excel.build_template()
    for name in excel.NAMES:
        ws, first = excel._anchor(wb, name)
        assert ws[first] is not None


def test_payload_shapes_match_named_ranges():
    from comps.pipeline import analyze
    a = analyze("AMZN", provider=SampleProvider())
    pl = excel.payload(a.result, a.candidates)
    assert len(pl["CompsData"]) == 1 + excel.MAX_PEERS
    assert all(len(r) == len(excel.INPUT_COLUMNS) for r in pl["CompsData"])
    assert len(pl["CompsMultiples"]) == excel.MAX_MULTIPLES
    assert len(pl["CompsScreen"]) == excel.MAX_SCREEN


def test_xlwings_refresh_with_fake_book():
    pytest.importorskip("xlwings")
    from comps import xl

    wb = excel.build_template()
    wb["Dashboard"]["C4"] = "Alibaba"
    wb["Dashboard"]["C5"] = 4
    xl.refresh(FakeBook(wb), provider=SampleProvider())
    assert wb["Dashboard"]["B9"].value.startswith("✔ 'Alibaba' → BABA")
    assert wb["Comps"]["A6"].value == "BABA"
    assert wb["Comps"]["A10"].value and wb["Comps"]["A11"].value is None  # 4 peers, rest cleared

    # Second refresh with fewer peers must clear stale rows
    wb["Dashboard"]["C4"] = "AMZN"
    wb["Dashboard"]["C6"] = "MSFT, GOOGL"
    xl.refresh(FakeBook(wb), provider=SampleProvider())
    assert [wb["Comps"][f"A{r}"].value for r in (6, 7, 8, 9)] == ["AMZN", "MSFT", "GOOGL", None]


def test_xlwings_refresh_reports_errors_in_sheet():
    pytest.importorskip("xlwings")
    from comps import xl

    wb = excel.build_template()
    wb["Dashboard"]["C4"] = "Not A Real Company"
    xl.refresh(FakeBook(wb), provider=SampleProvider())
    assert wb["Dashboard"]["B9"].value.startswith("✖")


def test_multiple_key_aliases():
    pytest.importorskip("xlwings")
    from comps.xl import multiple_key

    assert multiple_key("EV / EBITDA") == "ev_ebitda"
    assert multiple_key("ev/ebitda") == "ev_ebitda"
    assert multiple_key("Fwd P/E") == "pe_fwd"
    assert multiple_key("P/E (LTM)") == "pe_ltm"
    with pytest.raises(ValueError):
        multiple_key("EV/Unicorns")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_workbook_recalculates_without_errors(tmp_path):
    recalc = next(iter(RECALC), None)
    if recalc is None:
        pytest.skip("recalc helper not available")
    path = tmp_path / "m.xlsx"
    excel.write_model_file(str(path))
    out = subprocess.run([sys.executable, str(recalc), str(path), "90"], capture_output=True, text=True)
    assert '"total_errors": 0' in out.stdout, out.stdout
    ws = load_workbook(path, data_only=True)["Dashboard"]
    assert ws[f"G{excel.D_BLEND}"].value == pytest.approx(130.46, abs=0.01)
