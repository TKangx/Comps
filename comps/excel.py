"""The CompsModel Excel workbook.

The workbook is a *template*: every multiple, peer statistic and implied price is
an Excel formula. Python only ever writes raw data into a handful of named ranges
(see ``payload``), so the exact same template is used three ways:

* ``CompsModel.xlsx`` + xlwings: type a company in the yellow cell, hit Run main
  (``comps/xl.py`` writes the payload through xlwings),
* the Streamlit "Download Excel" button and ``python -m comps`` (written with openpyxl),
* by hand: overwrite any blue input and everything recalculates.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .analysis import CompsResult
from .multiples import MULTIPLE_REASONS, MULTIPLES
from .peers import PeerScore

# ----------------------------------------------------------------------------- style
FONT = "Arial"
BLUE, BLACK, GREEN, WHITE, NAVY, GREY = "0000FF", "000000", "008000", "FFFFFF", "1F3864", "595959"
FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_HEADER_SEL = PatternFill("solid", fgColor="2E75B6")
FILL_TARGET = PatternFill("solid", fgColor="FFF2CC")
FILL_STATS = PatternFill("solid", fgColor="F2F2F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")

FMT_MM = '#,##0;(#,##0);"-"'
FMT_PRICE = '#,##0.00;(#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_X = '0.0"x";(0.0"x");"-"'
FMT_SIGNED = '+0.0%;-0.0%;0.0%'

# ---------------------------------------------------------------------------- layout
MAX_PEERS = 8
MAX_MULTIPLES = 6
MAX_SCREEN = 25

# Comps sheet input block (Python writes these; A..T)
INPUT_COLUMNS = [
    ("ticker", "Ticker", None, None),
    ("name", "Company", None, None),
    ("price", "Share Price", FMT_PRICE, 1),
    ("market_cap", "Market Cap ($mm)", FMT_MM, 1e-6),
    ("total_debt", "Total Debt ($mm)", FMT_MM, 1e-6),
    ("cash", "Cash ($mm)", FMT_MM, 1e-6),
    ("revenue", "LTM Revenue ($mm)", FMT_MM, 1e-6),
    ("ntm_revenue", "NTM Revenue ($mm)", FMT_MM, 1e-6),
    ("gross_profit", "LTM Gross Profit ($mm)", FMT_MM, 1e-6),
    ("ebitda", "LTM EBITDA ($mm)", FMT_MM, 1e-6),
    ("ebit", "LTM EBIT ($mm)", FMT_MM, 1e-6),
    ("net_income", "LTM Net Income ($mm)", FMT_MM, 1e-6),
    ("forward_earnings", "Forward Earnings ($mm)", FMT_MM, 1e-6),
    ("operating_cash_flow", "LTM Op. Cash Flow ($mm)", FMT_MM, 1e-6),
    ("free_cash_flow", "LTM FCF ($mm)", FMT_MM, 1e-6),
    ("book_value", "Book Value ($mm)", FMT_MM, 1e-6),
    ("revenue_growth", "Revenue Growth", FMT_PCT, 1),
    ("gross_margin", "Gross Margin", FMT_PCT, 1),
    ("ebitda_margin", "EBITDA Margin", FMT_PCT, 1),
    ("net_margin", "Net Margin", FMT_PCT, 1),
]
COL = {attr: i for i, (attr, *_rest) in enumerate(INPUT_COLUMNS, start=1)}
HEADER_OF = {attr: h for attr, h, _, _ in INPUT_COLUMNS}
L = {attr: get_column_letter(i) for attr, i in COL.items()}
LAST_INPUT = get_column_letter(len(INPUT_COLUMNS))            # T
EV_COL = len(INPUT_COLUMNS) + 2                                # V (U is a spacer)
L["ev"] = get_column_letter(EV_COL)
MULT_COL = {k: EV_COL + 1 + i for i, k in enumerate(MULTIPLES)}  # W..AF
ML = {k: get_column_letter(c) for k, c in MULT_COL.items()}
FIRST_MULT, LAST_MULT = ML[next(iter(MULTIPLES))], ML[list(MULTIPLES)[-1]]

C_CAP, C_HDR, C_TGT = 4, 5, 6
C_PEER1, C_PEERN = C_TGT + 1, C_TGT + MAX_PEERS                # 7..14
C_STAT1 = C_PEERN + 2                                          # 16
STATS = [("Maximum", "MAX({r})"), ("75th Percentile", "PERCENTILE({r},0.75)"), ("Mean", "AVERAGE({r})"),
         ("Median", "MEDIAN({r})"), ("25th Percentile", "PERCENTILE({r},0.25)"), ("Minimum", "MIN({r})")]
STAT_ROW = {name: C_STAT1 + i for i, (name, _) in enumerate(STATS)}
C_PREM = C_STAT1 + len(STATS) + 1

# Dashboard
D_MULT1 = 21
D_MULTN = D_MULT1 + MAX_MULTIPLES - 1                          # 26
D_BLEND = D_MULTN + 2                                          # 28
D_PEER_HDR = 37
D_PEER1 = D_PEER_HDR + 1
# Valuation
V_ROW1 = 9
V_BLEND = V_ROW1 + MAX_MULTIPLES + 1
# Lists
LISTS_LAST = 1 + len(MULTIPLES)

# Named ranges: inputs Python reads, anchors Python writes
NAMES = {
    "CompsInput": "Dashboard!$C$4",
    "CompsNumPeers": "Dashboard!$C$5",
    "CompsCustomPeers": "Dashboard!$C$6",
    "CompsStatus": "Dashboard!$B$9",
    "CompsSubtitle": "Dashboard!$B$12",
    "CompsProfile": "Dashboard!$C$17",
    "CompsProfileWhy": "Dashboard!$B$18",
    "CompsMultiples": f"Dashboard!$B${D_MULT1}:$C${D_MULTN}",
    "CompsNotes": "Dashboard!$B$33",
    "CompsPeerTable": f"Dashboard!$B${D_PEER1}:$G${D_PEER1 + MAX_PEERS - 1}",
    "CompsData": f"Comps!$A${C_TGT}:${LAST_INPUT}${C_PEERN}",
    "CompsScreen": f"'Peer Screen'!$A$5:$N${4 + MAX_SCREEN}",
}


def _style(cell, *, bold=False, color=BLACK, fmt=None, fill=None, align=None, size=10, italic=False):
    cell.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align == "wrap":
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    elif align:
        cell.alignment = Alignment(horizontal=align, vertical="center")
    return cell


def _put(ws, row, col, value, **kw):
    return _style(ws.cell(row=row, column=col, value=value), **kw)


def _header(ws, row, headers, start_col=1, fills=None):
    for i, h in enumerate(headers):
        c = _put(ws, row, start_col + i, h, bold=True, color=WHITE,
                 fill=(fills or {}).get(i, FILL_HEADER))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ========================================================================= template
def build_template() -> Workbook:
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    comps = wb.create_sheet("Comps")
    val = wb.create_sheet("Valuation")
    screen = wb.create_sheet("Peer Screen")
    lists = wb.create_sheet("Lists")
    notes = wb.create_sheet("Notes")

    _build_lists(lists)
    _build_comps(comps)
    _build_valuation(val)
    _build_dashboard(dash)
    _build_screen(screen)
    _build_notes(notes)

    for name, ref in NAMES.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    return wb


def _build_lists(ws):
    _header(ws, 1, ["Multiple", "Target metric (Comps header)", "Type", "Why use it", "NM cap"])
    for r, (key, m) in enumerate(MULTIPLES.items(), start=2):
        _put(ws, r, 1, m.label, bold=True)
        _put(ws, r, 2, HEADER_OF[m.metric])
        _put(ws, r, 3, "EV" if m.numerator == "ev" else "Equity")
        _put(ws, r, 4, MULTIPLE_REASONS[key])
        _put(ws, r, 5, f"=Comps!{ML[key]}${C_CAP}", color=GREEN, fmt=FMT_X)
    _put(ws, LISTS_LAST + 2, 1, "Lookup table used by the Dashboard dropdowns and the Valuation sheet. "
         "Edit the 'Why use it' text freely; keep the first three columns as they are.", italic=True, size=8)
    _widths(ws, {"A": 26, "B": 28, "C": 9, "D": 80, "E": 9})


def _build_comps(ws):
    _put(ws, 1, 1, "Comparable Companies Analysis", bold=True, size=14, color=NAVY)
    _put(ws, 2, 1, "$ in millions except per share (reporting currency converted to trading currency). "
         "Blue = input written by Refresh (safe to overwrite), black = formula. Highlighted headers = "
         "multiples picked on the Dashboard.", italic=True, size=9)

    _header(ws, C_HDR, [h for _, h, _, _ in INPUT_COLUMNS])
    _header(ws, C_HDR, ["Enterprise Value ($mm)"] + [m.label for m in MULTIPLES.values()], start_col=EV_COL)
    _put(ws, C_CAP, EV_COL, "NM if above →", italic=True, size=8, align="right")
    for key, m in MULTIPLES.items():
        c = _put(ws, C_CAP, MULT_COL[key], m.max_meaningful, color=BLUE, fmt=FMT_X, size=8, align="center")
        c.comment = Comment("Multiples above this cap (or with a negative denominator) show as NM and "
                            "are excluded from peer statistics. Edit to change.", "Comps")

    mcap, debt, cash = L["market_cap"], L["total_debt"], L["cash"]
    for r in range(C_TGT, C_PEERN + 1):
        tgt = r == C_TGT
        fill = FILL_TARGET if tgt else None
        for attr, _, fmt, _ in INPUT_COLUMNS:
            _put(ws, r, COL[attr], None, color=BLUE, fmt=fmt, fill=fill, bold=tgt)
        _put(ws, r, EV_COL, f'=IF($A{r}="","",IF(ISNUMBER({mcap}{r}),{mcap}{r}+N({debt}{r})-N({cash}{r}),"NM"))',
             fmt=FMT_MM, fill=fill, bold=tgt, align="right")
        for key, m in MULTIPLES.items():
            num = f"{L['ev']}{r}" if m.numerator == "ev" else f"{mcap}{r}"
            den = f"{L[m.metric]}{r}"
            cap = f"{ML[key]}${C_CAP}"
            f = (f'=IF($A{r}="","",IF(AND(ISNUMBER({den}),ISNUMBER({num})),IF(AND({den}>0,{num}>0),'
                 f'IF({num}/{den}<={cap},{num}/{den},"NM"),"NM"),"NM"))')
            _put(ws, r, MULT_COL[key], f, fmt=FMT_X, fill=fill, bold=tgt, align="right")

    _put(ws, C_STAT1 - 1, 2, "Peer statistics (excluding target)", bold=True, italic=True, size=9)
    stat_cols = [a for a, _, fmt, _ in INPUT_COLUMNS if fmt == FMT_PCT] + list(MULTIPLES)
    for label, tmpl in STATS:
        r = STAT_ROW[label]
        _put(ws, r, 2, label, bold=True, fill=FILL_STATS)
        for attr in stat_cols:
            col = MULT_COL.get(attr) or COL[attr]
            letter = get_column_letter(col)
            rng = f"{letter}{C_PEER1}:{letter}{C_PEERN}"
            _put(ws, r, col, f'=IF(COUNT({rng})=0,"NM",{tmpl.format(r=rng)})',
                 fmt=FMT_X if attr in MULTIPLES else FMT_PCT, fill=FILL_STATS, align="right")

    _put(ws, C_PREM, 2, '=IF($A$6="","Target",$A$6)&" vs. peer median"', bold=True, fill=FILL_TARGET)
    med = STAT_ROW["Median"]
    for attr in stat_cols:
        col = MULT_COL.get(attr) or COL[attr]
        letter = get_column_letter(col)
        t, m = f"{letter}{C_TGT}", f"{letter}{med}"
        expr = f"{t}/{m}-1" if attr in MULTIPLES else f"{t}-{m}"
        _put(ws, C_PREM, col, f'=IF(AND(ISNUMBER({t}),ISNUMBER({m})),{expr},"NM")',
             fmt=FMT_SIGNED, fill=FILL_TARGET, bold=True, align="right")
    _put(ws, C_PREM + 1, 2, "Margin/growth columns: difference in percentage points. Multiple columns: "
         "premium (+) / discount (−) vs. the peer median.", italic=True, size=8)

    # highlight the multiples currently picked on the Dashboard
    pick = f"Dashboard!$B${D_MULT1}:$B${D_MULTN}"
    ws.conditional_formatting.add(
        f"{FIRST_MULT}{C_HDR}:{LAST_MULT}{C_HDR}",
        FormulaRule(formula=[f'COUNTIF({pick},{FIRST_MULT}{C_HDR})>0'], fill=FILL_HEADER_SEL))

    _widths(ws, {"A": 9, "B": 30, get_column_letter(EV_COL - 1): 3})
    for i in list(range(3, len(INPUT_COLUMNS) + 1)) + list(range(EV_COL, EV_COL + 1 + len(MULTIPLES))):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = ws.cell(row=C_TGT, column=3)


def _build_valuation(ws):
    _put(ws, 1, 1, "Implied Valuation from Peer Multiples", bold=True, size=14, color=NAVY)
    _put(ws, 2, 1, "Implied price = (peer multiple × target metric − net debt for EV multiples) ÷ shares. "
         "Multiples and blend flags come from the Dashboard.", italic=True, size=9)
    t = C_TGT
    _put(ws, 4, 1, "Current share price", bold=True)
    _put(ws, 4, 3, f"=Comps!{L['price']}{t}", color=GREEN, fmt=FMT_PRICE)
    _put(ws, 5, 1, "Shares outstanding (mm, implied)", bold=True)
    _put(ws, 5, 3, f'=IFERROR(Comps!{L["market_cap"]}{t}/Comps!{L["price"]}{t},"NM")', fmt=FMT_MM)
    _put(ws, 6, 1, "Net debt ($mm)", bold=True)
    _put(ws, 6, 3, f"=N(Comps!{L['total_debt']}{t})-N(Comps!{L['cash']}{t})", fmt=FMT_MM)

    _header(ws, V_ROW1 - 1, ["Multiple", "Target metric", "Type", "Metric ($mm)", "Target multiple",
                             "Peer 25th", "Peer median", "Peer 75th", "Implied @ 25th", "Implied @ median",
                             "Implied @ 75th", "Upside @ median", "Use in blend"])
    lab = f"Lists!$A$2:$A${LISTS_LAST}"
    mh, mr = f"Comps!${FIRST_MULT}${C_HDR}:${LAST_MULT}${C_HDR}", lambda row: f"Comps!${FIRST_MULT}${row}:${LAST_MULT}${row}"
    for i in range(MAX_MULTIPLES):
        r = V_ROW1 + i
        d = D_MULT1 + i
        blank = f'$A{r}=""'
        _put(ws, r, 1, f'=IF(Dashboard!$B{d}="","",Dashboard!$B{d})', color=GREEN, bold=True)
        _put(ws, r, 2, f'=IF({blank},"",IFERROR(INDEX(Lists!$B$2:$B${LISTS_LAST},MATCH($A{r},{lab},0)),"?"))')
        _put(ws, r, 3, f'=IF({blank},"",IFERROR(INDEX(Lists!$C$2:$C${LISTS_LAST},MATCH($A{r},{lab},0)),"?"))',
             align="center")
        _put(ws, r, 4, f'=IF({blank},"",IFERROR(INDEX(Comps!$A${t}:${LAST_INPUT}${t},'
                       f'MATCH($B{r},Comps!$A${C_HDR}:${LAST_INPUT}${C_HDR},0)),"NM"))', fmt=FMT_MM, align="right")
        for col, row in ((5, t), (6, STAT_ROW["25th Percentile"]), (7, STAT_ROW["Median"]),
                         (8, STAT_ROW["75th Percentile"])):
            _put(ws, r, col, f'=IF({blank},"",IFERROR(INDEX({mr(row)},MATCH($A{r},{mh},0)),"NM"))',
                 fmt=FMT_X, align="right")
        for col, src in ((9, "F"), (10, "G"), (11, "H")):
            _put(ws, r, col, f'=IF({blank},"",IF(AND(ISNUMBER({src}{r}),ISNUMBER($D{r}),ISNUMBER($C$5)),'
                             f'IF($D{r}>0,MAX(0,({src}{r}*$D{r}-IF($C{r}="EV",$C$6,0))/$C$5),"NM"),"NM"))',
                 fmt=FMT_PRICE, align="right")
        _put(ws, r, 12, f'=IF({blank},"",IF(ISNUMBER(J{r}),J{r}/$C$4-1,"NM"))', fmt=FMT_SIGNED, align="right")
        _put(ws, r, 13, f'=IF({blank},"",Dashboard!$C{d})', color=GREEN, align="center")
    last = V_ROW1 + MAX_MULTIPLES - 1
    _put(ws, V_BLEND, 1, "Blended implied price (avg. of 'Y' rows)", bold=True, fill=FILL_TARGET)
    for col, src in ((9, "I"), (10, "J"), (11, "K")):
        _put(ws, V_BLEND, col, f'=IFERROR(AVERAGEIFS({src}{V_ROW1}:{src}{last},$M${V_ROW1}:$M${last},"Y"),"NM")',
             fmt=FMT_PRICE, bold=True, fill=FILL_TARGET, align="right")
    _put(ws, V_BLEND, 12, f'=IF(ISNUMBER(J{V_BLEND}),J{V_BLEND}/$C$4-1,"NM")', fmt=FMT_SIGNED, bold=True,
         fill=FILL_TARGET, align="right")
    ws.conditional_formatting.add(f"L{V_ROW1}:L{V_BLEND}", FormulaRule(
        formula=[f"AND(ISNUMBER(L{V_ROW1}),L{V_ROW1}<0)"], font=Font(name=FONT, color="C00000")))
    _widths(ws, {"A": 36, "B": 24, "C": 8, "D": 13, "E": 12, "F": 11, "G": 11, "H": 11, "I": 13, "J": 13,
                 "K": 13, "L": 13, "M": 10})


def _build_dashboard(ws):
    t = C_TGT
    _put(ws, 2, 2, "SMIF Comps Analyzer", bold=True, size=16, color=NAVY)

    _put(ws, 4, 2, "Company name or ticker", bold=True)
    _put(ws, 4, 3, "AMZN", color=BLUE, bold=True, fill=FILL_INPUT, size=12)
    _put(ws, 5, 2, "Number of peers (3–8)", bold=True)
    _put(ws, 5, 3, 5, color=BLUE, fill=FILL_INPUT, align="left")
    _put(ws, 6, 2, "Custom peers (optional)", bold=True)
    _put(ws, 6, 3, None, color=BLUE, fill=FILL_INPUT)
    _put(ws, 4, 5, "① Type a company or ticker in the yellow cell", size=9, color=GREY)
    _put(ws, 5, 5, "② Click the xlwings tab → Run main (or your Refresh button)", size=9, color=GREY)
    _put(ws, 6, 5, "③ Leave custom peers blank to auto-pick, or list tickers: MSFT, GOOGL", size=9, color=GREY)
    dv = DataValidation(type="whole", operator="between", formula1="3", formula2=str(MAX_PEERS), showErrorMessage=True,
                        errorTitle="Peers", error=f"Enter a whole number from 3 to {MAX_PEERS}.")
    ws.add_data_validation(dv)
    dv.add("C5")

    _put(ws, 9, 2, "Status: not refreshed yet.", italic=True, size=9, color=GREY)
    ws.merge_cells("B9:I9")

    _put(ws, 11, 2, f'=IF(Comps!$A${t}="","",Comps!$B${t}&" ("&Comps!$A${t}&")")', bold=True, size=14, color=NAVY)
    _put(ws, 12, 2, None, size=9, color=GREY)
    ws.merge_cells("B12:I12")

    stats = [("Share price", f"=Comps!{L['price']}{t}", FMT_PRICE),
             ("Market cap ($mm)", f"=Comps!{L['market_cap']}{t}", FMT_MM),
             ("Enterprise value ($mm)", f"=Comps!{L['ev']}{t}", FMT_MM),
             ("Revenue growth", f"=Comps!{L['revenue_growth']}{t}", FMT_PCT),
             ("EBITDA margin", f"=Comps!{L['ebitda_margin']}{t}", FMT_PCT),
             ("Net margin", f"=Comps!{L['net_margin']}{t}", FMT_PCT)]
    for i, (lab, f, fmt) in enumerate(stats):
        _put(ws, 14, 2 + i, lab, size=8, color=GREY)
        _put(ws, 15, 2 + i, f, color=GREEN, fmt=fmt, bold=True, size=12, align="left")

    _put(ws, 17, 2, "Valuation profile", bold=True)
    _put(ws, 17, 3, None, bold=True, color=NAVY)
    _put(ws, 18, 2, None, align="wrap", size=9)
    ws.merge_cells("B18:I18")
    ws.row_dimensions[18].height = 28

    _header(ws, D_MULT1 - 1, ["Multiple (pick from list)", "Use in blend", "Target", "Peer median",
                              "Premium / (discount)", "Implied price @ median", "Upside", "Why this multiple"],
            start_col=2)
    for i in range(MAX_MULTIPLES):
        r, v = D_MULT1 + i, V_ROW1 + i
        blank = f'$B{r}=""'
        _put(ws, r, 2, None, color=BLUE, bold=True, fill=FILL_INPUT)
        _put(ws, r, 3, None, color=BLUE, fill=FILL_INPUT, align="center")
        _put(ws, r, 4, f'=IF({blank},"",Valuation!E{v})', color=GREEN, fmt=FMT_X, align="right")
        _put(ws, r, 5, f'=IF({blank},"",Valuation!G{v})', color=GREEN, fmt=FMT_X, align="right")
        _put(ws, r, 6, f'=IF({blank},"",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),D{r}/E{r}-1,"NM"))',
             fmt=FMT_SIGNED, align="right")
        _put(ws, r, 7, f'=IF({blank},"",Valuation!J{v})', color=GREEN, fmt=FMT_PRICE, align="right")
        _put(ws, r, 8, f'=IF({blank},"",Valuation!L{v})', color=GREEN, fmt=FMT_SIGNED, align="right")
        _put(ws, r, 9, f'=IF({blank},"",IFERROR(INDEX(Lists!$D$2:$D${LISTS_LAST},'
                       f'MATCH($B{r},Lists!$A$2:$A${LISTS_LAST},0)),""))', size=9)
    mv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${LISTS_LAST}", allow_blank=True)
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(mv)
    ws.add_data_validation(yn)
    mv.add(f"B{D_MULT1}:B{D_MULTN}")
    yn.add(f"C{D_MULT1}:C{D_MULTN}")
    for col in "FH":
        ws.conditional_formatting.add(f"{col}{D_MULT1}:{col}{D_BLEND}", FormulaRule(
            formula=[f"AND(ISNUMBER({col}{D_MULT1}),{col}{D_MULT1}<0)"], font=Font(name=FONT, color="C00000")))

    _put(ws, D_BLEND, 2, "Blended implied price ('Y' multiples, peer median)", bold=True, fill=FILL_TARGET)
    for c in range(3, 7):
        _put(ws, D_BLEND, c, None, fill=FILL_TARGET)
    _put(ws, D_BLEND, 7, f"=Valuation!J{V_BLEND}", color=GREEN, fmt=FMT_PRICE, bold=True, fill=FILL_TARGET,
         align="right")
    _put(ws, D_BLEND, 8, f"=Valuation!L{V_BLEND}", color=GREEN, fmt=FMT_SIGNED, bold=True, fill=FILL_TARGET,
         align="right")
    price = f"Comps!$C${t}"
    _put(ws, 30, 2, f'=IF(OR(Comps!$A${t}="",NOT(ISNUMBER(G{D_BLEND}))),"Refresh to see the summary.",'
                    f'Comps!$A${t}&" trades at "&TEXT({price},"#,##0.00")&" vs. "&TEXT(G{D_BLEND},"#,##0.00")&'
                    f'" implied by peer-median multiples: a "&TEXT(ABS({price}/G{D_BLEND}-1),"0%")&'
                    f'IF({price}>=G{D_BLEND}," premium"," discount")&" to the blended peer-implied value.")',
         bold=True, align="wrap")
    ws.merge_cells("B30:I30")
    ws.row_dimensions[30].height = 30

    _put(ws, 32, 2, "Analyst notes (written at refresh)", bold=True, color=NAVY)
    _put(ws, 33, 2, None, align="wrap", size=9)
    ws.merge_cells("B33:I35")

    _header(ws, D_PEER_HDR, ["Peer company", "Ticker", "Industry", "Market cap ($mm)", "Similarity (0–100)",
                             "Why it's a comp", "", ""], start_col=2)
    ws.merge_cells(start_row=D_PEER_HDR, start_column=7, end_row=D_PEER_HDR, end_column=9)
    for i in range(MAX_PEERS):
        r = D_PEER1 + i
        for c, fmt in zip(range(2, 8), (None, None, None, FMT_MM, "0", None)):
            _put(ws, r, c, None, color=BLUE, fmt=fmt, size=9 if c == 7 else 10, bold=c == 3,
                 align="center" if c in (3, 6) else None)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
    _put(ws, D_PEER1 + MAX_PEERS + 1, 2, "Blue = written by Refresh · Yellow = your inputs · Green = links · "
         "Black = formulas. Data: Yahoo Finance – verify key figures against filings before presenting.",
         italic=True, size=8, color=GREY)
    _widths(ws, {"A": 2, "B": 34, "C": 16, "D": 22, "E": 16, "F": 18, "G": 18, "H": 12, "I": 70})


def _build_screen(ws):
    _put(ws, 1, 1, "Peer Candidate Screen", bold=True, size=14, color=NAVY)
    _put(ws, 2, 1, "Every candidate considered, ranked by similarity (0–100). Weights: industry 35, investor "
         "grouping 10, size 20, revenue growth 15, EBITDA margin 10, gross margin 10.", italic=True, size=9)
    _header(ws, 4, ["Rank", "Ticker", "Company", "Industry", "Market cap ($mm)", "Rev growth", "EBITDA margin",
                    "Industry pts", "Investor pts", "Size pts", "Growth pts", "EBITDA mgn pts", "Gross mgn pts",
                    "Total score"])
    fmts = [None, None, None, None, FMT_MM, FMT_PCT, FMT_PCT] + ["0.0"] * 7
    for r in range(5, 5 + MAX_SCREEN):
        for c, f in enumerate(fmts, start=1):
            _put(ws, r, c, None, color=BLUE, fmt=f, bold=c in (2, 14))
    _widths(ws, dict(zip("ABCDEFGHIJKLMN", (6, 9, 30, 26, 14, 11, 11, 10, 10, 10, 10, 11, 11, 11))))
    ws.freeze_panes = "C5"


def _build_notes(ws):
    rows = [
        ("How to use", None),
        ("1", "Type a company name ('Amazon') or ticker ('AMZN') into the yellow cell on the Dashboard."),
        ("2", "Click the xlwings ribbon tab → Run main. Python pulls Yahoo Finance data, picks peers and "
              "multiples, and writes them into the blue cells. Everything else is formulas."),
        ("3", "Change the multiples with the dropdowns and the Y/N blend flags on the Dashboard – no "
              "refresh needed. Overwrite any blue input on the Comps sheet to test your own numbers."),
        ("Keep the name", "Run main calls the Python file with the same name as the workbook (CompsModel.py). "
                          "To keep a snapshot of a company, use File → Save a Copy."),
        ("", None),
        ("Methodology", None),
        ("Peers", "Candidates come from Yahoo's industry constituents, Yahoo's 'people also watch' list and "
                  "(if thin) sector leaders, scored on industry, investor grouping, size, growth and margins."),
        ("Multiples", "Picked by profile (financial, REIT, unprofitable, high-growth, capital-intensive, mature). "
                      "EV = market cap + total debt − cash."),
        ("NM", "Not meaningful: negative/zero denominator or above the cap in row 4 of Comps. Excluded from stats."),
        ("NTM Revenue", "Time-weighted blend of consensus current-FY and next-FY revenue."),
        ("Fwd Earnings", "Market cap ÷ Yahoo forward P/E (consensus next fiscal year EPS)."),
        ("EBIT", "LTM revenue × LTM operating margin."),
        ("Currency", "ADR financials are converted to the trading currency at spot FX."),
        ("Limitations", "Yahoo data can lag filings and is not adjusted for one-offs, leases, minority interest "
                        "or preferreds. Sanity-check key inputs against the 10-K/10-Q before presenting."),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        _put(ws, i, 1, k, bold=True, color=NAVY if v is None else BLACK)
        if v:
            _put(ws, i, 2, v, align="wrap")
    _widths(ws, {"A": 16, "B": 110})


# ========================================================================== payload
def _peer_reason(s: Optional[PeerScore]) -> str:
    return "; ".join(s.reasons) if s else ""


def payload(res: CompsResult, candidates: Optional[list[PeerScore]] = None, query: str = "",
            status: Optional[str] = None) -> dict[str, list[list]]:
    """Every value Python writes, keyed by the named range it goes into."""
    t = res.target
    peers = list(res.peers)[:MAX_PEERS]
    score_of = {s.ticker: s for s in (candidates or [])}

    def row(c):
        out = []
        for attr, _, _, scale in INPUT_COLUMNS:
            v = getattr(c, attr)
            out.append(v * scale if scale and v is not None else v)
        return out

    blank_inputs = [None] * len(INPUT_COLUMNS)
    data = [row(t)] + [row(p) for p in peers] + [blank_inputs] * (MAX_PEERS - len(peers))

    mults = [[MULTIPLES[k].label, "Y" if k in res.selection.primary else "N"]
             for k in res.selection.multiples[:MAX_MULTIPLES]]
    mults += [[None, None]] * (MAX_MULTIPLES - len(mults))

    peer_rows = [[p.name, p.ticker, p.industry, (p.market_cap or 0) * 1e-6,
                  score_of[p.ticker].score if p.ticker in score_of else None, _peer_reason(score_of.get(p.ticker))]
                 for p in peers]
    peer_rows += [[None] * 6] * (MAX_PEERS - len(peer_rows))

    screen = []
    for i, s in enumerate((candidates or [])[:MAX_SCREEN], start=1):
        c = s.company
        screen.append([i, c.ticker, c.name, c.industry, (c.market_cap or 0) * 1e-6, c.revenue_growth,
                       c.ebitda_margin] + [round(s.components.get(k, 0.0), 1) for k in
                                           ("industry", "market", "size", "growth", "ebitda_margin",
                                            "gross_margin")] + [s.score])
    screen += [[None] * 14] * (MAX_SCREEN - len(screen))

    if status is None:
        resolved = f"'{query}' → {t.ticker} · " if query and query.strip().upper() != t.ticker else ""
        status = (f"✔ {resolved}{t.name} · {len(peers)} peers · {t.currency} · data as of {t.as_of} · "
                  f"refreshed {datetime.now():%Y-%m-%d %H:%M}")
    return {
        "CompsStatus": [[status]],
        "CompsSubtitle": [[" · ".join(x for x in (t.sector, t.industry, t.country, t.currency) if x)]],
        "CompsProfile": [[res.selection.profile.replace("_", " ").title()]],
        "CompsProfileWhy": [[res.selection.explanation]],
        "CompsMultiples": mults,
        "CompsNotes": [[res.verdict]],
        "CompsPeerTable": peer_rows,
        "CompsData": data,
        "CompsScreen": screen,
    }


def _anchor(wb: Workbook, name: str):
    sheet, ref = next(iter(wb.defined_names[name].destinations))
    first = ref.split(":")[0].replace("$", "")
    return wb[sheet], first


def fill(wb: Workbook, values: dict[str, list[list]]) -> Workbook:
    """Write a payload into an openpyxl workbook (xlwings has its own writer in comps/xl.py)."""
    for name, block in values.items():
        ws, first = _anchor(wb, name)
        top = ws[first]
        for i, r in enumerate(block):
            for j, v in enumerate(r):
                ws.cell(row=top.row + i, column=top.column + j).value = v
    return wb


def build_workbook(res: CompsResult, candidates: Optional[list[PeerScore]] = None, query: str = "",
                   status: Optional[str] = None) -> Workbook:
    wb = build_template()
    fill(wb, payload(res, candidates, query=query, status=status))
    wb["Dashboard"]["C4"].value = query or res.target.ticker
    wb["Dashboard"]["C5"].value = max(3, min(MAX_PEERS, len(res.peers)))
    return wb


def to_bytes(res: CompsResult, candidates: Optional[list[PeerScore]] = None, query: str = "") -> bytes:
    buf = BytesIO()
    build_workbook(res, candidates, query=query).save(buf)
    return buf.getvalue()


def write_model_file(path: str = "CompsModel.xlsx") -> None:
    """Regenerate the CompsModel.xlsx template, pre-filled with offline sample data."""
    from .pipeline import analyze
    from .sample import SampleProvider

    a = analyze("AMZN", provider=SampleProvider())
    wb = build_workbook(a.result, a.candidates, query="AMZN",
                        status="SAMPLE DATA (illustrative, not live) – type a company in the yellow cell, "
                               "then click xlwings ▸ Run main to load live Yahoo Finance data.")
    wb.save(path)


if __name__ == "__main__":
    import sys

    out = sys.argv[1] if len(sys.argv) > 1 else "CompsModel.xlsx"
    write_model_file(out)
    print(f"Wrote {out}")
